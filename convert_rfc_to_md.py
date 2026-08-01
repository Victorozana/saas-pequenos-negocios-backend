from pathlib import Path
from docx import Document

SOURCE = Path('docs/rfcs/RFC-Arquitetura-Sistema-Agendamento-Restaurantes.docx')
TARGET = Path('docs/rfcs/RFC-Arquitetura-Sistema-Agendamento-Restaurantes.md')


def escape_cell(value: str) -> str:
    return ' '.join(value.replace('|', r'\|').splitlines()).strip()


def render_runs(paragraph) -> str:
    chunks = []
    for run in paragraph.runs:
        text = run.text
        if not text:
            continue
        visible = text.rstrip()
        trailing = text[len(visible):]
        if run.bold and run.italic:
            text = f'***{visible}***{trailing}'
        elif run.bold:
            text = f'**{visible}**{trailing}'
        elif run.italic:
            text = f'*{visible}*{trailing}'
        chunks.append(text)
    return ''.join(chunks).strip()


document = Document(SOURCE)
output = [
    '<!--',
    'Fonte: RFC-Arquitetura-Sistema-Agendamento-Restaurantes.docx',
    'Formato: Markdown para consumo por pessoas e IAs.',
    '-->',
    ''
]

table_index = 0
first_content = True
body = list(document.element.body)
paragraph_map = {paragraph._p: paragraph for paragraph in document.paragraphs}
table_map = {table._tbl: table for table in document.tables}

for element in body:
    if element in paragraph_map:
        paragraph = paragraph_map[element]
        text = render_runs(paragraph)
        if not text:
            continue
        style = paragraph.style.name
        if first_content:
            output.extend([f'# {text}', ''])
            first_content = False
        elif style.startswith('Heading '):
            level = style.split()[-1]
            output.extend([f"{'#' * int(level)} {text}", ''])
        elif style.startswith('List Bullet'):
            output.append(f'- {text}')
        elif style.startswith('List Number'):
            output.append(f'1. {text}')
        else:
            output.extend([text, ''])
    elif element in table_map:
        table = table_map[element]
        table_index += 1
        rows = [[escape_cell(cell.text) for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        header = rows[0]
        output.append('| ' + ' | '.join(header) + ' |')
        output.append('| ' + ' | '.join('---' for _ in header) + ' |')
        for row in rows[1:]:
            row += [''] * (len(header) - len(row))
            output.append('| ' + ' | '.join(row[:len(header)]) + ' |')
        output.append('')

text = '\n'.join(output).replace('\n\n\n', '\n\n').strip() + '\n'
TARGET.write_text(text, encoding='utf-8')
print(f'{TARGET} ({len(text.splitlines())} linhas)')
